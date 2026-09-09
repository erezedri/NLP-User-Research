"""
excel_builder.py — Write classified & tagged reviews to a styled Excel file.

Schema:
    Column A: DOCTOR/PATIENT   (persona — "Patient" unless doctor reviews collected)
    Column B: ORIGINAL REVIEW  (raw review text)
    Column C: CLASSIFICATION   ("Good Experience" / "Bad Experience")
    Column D: TAG              (category tag string)

Layout rules:
    - Row 1: bold header with background fill
    - Good rows: light green fill (#D5F5E3)
    - Bad rows:  light red fill   (#FADBD8)
    - Balanced dataset: equal Good and Bad count (trim the larger class)

Known pitfall: always close the file in Excel before calling write_excel(),
               or openpyxl will raise PermissionError.
"""

import logging
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

logger = logging.getLogger(__name__)

COLUMNS = ["DOCTOR/PATIENT", "ORIGINAL REVIEW", "CLASSIFICATION", "TAG"]

FILL_HEADER = PatternFill(fill_type="solid", fgColor="2E4057")
FILL_GOOD = PatternFill(fill_type="solid", fgColor="D5F5E3")
FILL_BAD = PatternFill(fill_type="solid", fgColor="FADBD8")

FONT_HEADER = Font(bold=True, color="FFFFFF", size=11)
FONT_BODY = Font(size=10)


def _balance(
    rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Trim the larger class so Good and Bad counts are equal."""
    good = [r for r in rows if r["classification"] == "Good Experience"]
    bad = [r for r in rows if r["classification"] == "Bad Experience"]
    target = min(len(good), len(bad))
    if len(good) != len(bad):
        logger.warning(
            "Balancing dataset: %d good / %d bad → trimming to %d each.",
            len(good),
            len(bad),
            target,
        )
    return good[:target] + bad[:target]


def write_excel(
    rows: list[dict[str, Any]],
    output_path: str | Path,
    balance: bool = True,
    persona_default: str = "Patient",
) -> Path:
    """
    Write reviews to a styled Excel file.

    Args:
        rows: list of dicts with keys: text, classification, tag, persona (optional)
        output_path: destination .xlsx file path
        balance: if True, trim to equal Good/Bad counts before writing
        persona_default: value for DOCTOR/PATIENT column when persona is absent

    Returns:
        Path to the written file.
    """
    if balance:
        rows = _balance(rows)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reviews"

    # Header row
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 24
    ws.row_dimensions[1].height = 20

    # Data rows
    for row_idx, review in enumerate(rows, start=2):
        classification = review.get("classification", "")
        fill = FILL_GOOD if classification == "Good Experience" else FILL_BAD
        values = [
            review.get("persona", persona_default),
            review.get("text", ""),
            classification,
            review.get("tag", ""),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.font = FONT_BODY
            cell.alignment = Alignment(wrap_text=(col_idx == 2), vertical="top")

    ws.freeze_panes = "A2"

    try:
        wb.save(output_path)
        logger.info("Excel written: %s (%d rows)", output_path, len(rows))
    except PermissionError:
        raise PermissionError(
            f"Cannot write to {output_path} — close the file in Excel first."
        )

    return output_path
